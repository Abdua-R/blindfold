import csv
import json
import sys
from pathlib import Path

import docx
import openpyxl
import pptx
import pypdf

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "eval"))
import make_formats  # noqa: E402
import make_synthetic  # noqa: E402

FILES = ("clients.xlsx", "clients.csv", "note.docx", "deck.pptx", "letter.pdf", "notes.md", "page.html")


def read_back(p: Path) -> str:
    if p.suffix == ".xlsx":
        ws = openpyxl.load_workbook(p).active
        return "\n".join("\t".join("" if v is None else str(v) for v in row) for row in ws.iter_rows(values_only=True))
    if p.suffix == ".csv":
        with p.open(encoding="utf-8", newline="") as f:
            return "\n".join("\t".join(r) for r in csv.reader(f))
    if p.suffix == ".docx":
        d = docx.Document(p)
        return "\n".join([x.text for x in d.paragraphs] +
                         [c.text for t in d.tables for r in t.rows for c in r.cells])
    if p.suffix == ".pptx":
        out = []
        for slide in pptx.Presentation(p).slides:
            for sh in slide.shapes:
                if sh.has_text_frame:
                    out.append(sh.text_frame.text)
                if sh.has_table:
                    out += [c.text for r in sh.table.rows for c in r.cells]
        return "\n".join(out)
    if p.suffix == ".pdf":
        return "\n".join(pg.extract_text() for pg in pypdf.PdfReader(p).pages)
    return p.read_text(encoding="utf-8")


def test_v11_format_fixtures_hold_their_seeded_values(tmp_path):
    make_formats.build(tmp_path)
    seeds = {d["file"]: d["items"] for d in json.loads((tmp_path / "seeds.json").read_text(encoding="utf-8"))}
    assert set(seeds) == set(FILES)
    for name in FILES:
        text = read_back(tmp_path / name)
        assert seeds[name], name
        for item in seeds[name]:
            assert item["value"] in text, (name, item)


def test_v11_formats_use_the_synthetic_clients(tmp_path):
    make_formats.build(tmp_path)
    make_synthetic.build(tmp_path / "syn", n_docs=8)
    syn = json.loads((tmp_path / "syn" / "seeds.json").read_text(encoding="utf-8"))
    xlsx = next(d for d in json.loads((tmp_path / "seeds.json").read_text(encoding="utf-8"))
                if d["file"] == "clients.xlsx")
    names = [i["value"] for i in xlsx["items"] if i["type"] == "Person"]
    assert names == [d["items"][0]["value"] for d in syn]
