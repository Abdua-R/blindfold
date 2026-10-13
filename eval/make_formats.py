"""Fake client files in common office formats, seeded with the same clients as make_synthetic."""
import csv
import json
import random
import sys
from pathlib import Path

import docx
import openpyxl
import pptx
from fpdf import FPDF
from pptx.util import Inches

from make_synthetic import client

FONT = r"C:\Windows\Fonts\arial.ttf"
HEAD = ["F.I.Sh.", "JShShIR", "hisob raqami", "karta", "telefon", "qarz summasi", "kechikish kunlari"]


def _ids(c, *keys):
    types = {"name": "Person", "first": "Person", "pinfl": "PINFL", "account": "Account", "card": "Card",
             "phone": "Phone"}
    return [{"type": types[k], "value": c[k]} for k in keys]


def build(out_dir: Path, seed: int = 7) -> list[dict]:
    rng = random.Random(seed)
    out_dir.mkdir(parents=True, exist_ok=True)
    cl = [client(rng, k % 3 == 1) for k in range(8)]   # k % 3 == 1: make_synthetic's Russian template
    rows = [[c["name"], c["pinfl"], c["account"], c["card"], c["phone"],
             rng.randint(5, 900) * 100_000, rng.randint(1, 180)] for c in cl]
    table_ids = [i for c in cl for i in _ids(c, "name", "pinfl", "account", "card", "phone")]
    seeds = [{"file": "clients.xlsx", "items": table_ids}, {"file": "clients.csv", "items": table_ids}]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mijozlar"
    for r in [HEAD] + rows:
        ws.append(r)
    wb.save(out_dir / "clients.xlsx")
    with (out_dir / "clients.csv").open("w", encoding="utf-8", newline="") as f:
        csv.writer(f).writerows([HEAD] + rows)

    a, b = cl[0], cl[2]
    d = docx.Document()
    d.add_heading("Mijoz bo'yicha xizmat yozuvi", 1)
    d.add_paragraph(f"Mijoz {a['name']} (JShShIR {a['pinfl']}) kredit bo'yicha 92 kun kechikkan. "
                    f"Hisob raqami {a['account']}, telefon {a['phone']}.")
    t = d.add_table(rows=2, cols=3)
    for j, v in enumerate(("Kafil", "JShShIR", "Telefon")):
        t.cell(0, j).text = v
    for j, v in enumerate((b["name"], b["pinfl"], b["phone"])):
        t.cell(1, j).text = v
    d.save(out_dir / "note.docx")
    seeds.append({"file": "note.docx", "items": _ids(a, "name", "pinfl", "account", "phone") +
                  _ids(b, "name", "pinfl", "phone")})

    c, e = cl[3], cl[4]
    deck = pptx.Presentation()

    def box(slide, text):
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text_frame.text = text

    s1 = deck.slides.add_slide(deck.slide_layouts[6])
    box(s1, f"Muammoli qarz: {c['name']}")
    s2 = deck.slides.add_slide(deck.slide_layouts[6])
    box(s2, f"JShShIR {c['pinfl']}, telefon {c['phone']}")
    s3 = deck.slides.add_slide(deck.slide_layouts[6])
    tbl = s3.shapes.add_table(2, 3, Inches(1), Inches(1), Inches(8), Inches(1)).table
    for j, v in enumerate(("F.I.Sh.", "JShShIR", "Telefon")):
        tbl.cell(0, j).text = v
    for j, v in enumerate((e["name"], e["pinfl"], e["phone"])):
        tbl.cell(1, j).text = v
    deck.save(out_dir / "deck.pptx")
    seeds.append({"file": "deck.pptx", "items": _ids(c, "name", "pinfl", "phone") + _ids(e, "name", "pinfl", "phone")})

    g = cl[5]
    pdf = FPDF()
    pdf.add_page()
    pdf.add_font("Arial", "", FONT)
    pdf.set_font("Arial", size=12)
    pdf.multi_cell(0, 8, f"Hurmatli {g['name']}!\n\n"
                         f"Sizning JShShIR {g['pinfl']} bo\u02bbyicha ochilgan {g['account']} hisob raqamidagi "
                         f"kredit qarzdorligi muddati o\u02bbtgan. Iltimos, filialga murojaat qiling.\n\n"
                         "\u0421 \u0443\u0432\u0430\u0436\u0435\u043d\u0438\u0435\u043c, "
                         "\u043e\u0442\u0434\u0435\u043b \u0440\u0430\u0431\u043e\u0442\u044b \u0441 "
                         "\u043a\u043b\u0438\u0435\u043d\u0442\u0430\u043c\u0438")
    pdf.output(str(out_dir / "letter.pdf"))
    seeds.append({"file": "letter.pdf", "items": _ids(g, "name", "pinfl", "account")})

    md = cl[6:8] + [cl[0]]
    lines = ["# Qo'ng'iroqlar rejasi", "", "| F.I.Sh. | Telefon | Qarz |", "|---|---|---|"]
    lines += [f"| {x['name']} | {x['phone']} | {r[5]} |" for x, r in zip(md, rows[6:8] + rows[:1])]
    lines += ["", f"{cl[2]['first']} ertaga qo'ng'iroq qilishini aytdi."]
    (out_dir / "notes.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    seeds.append({"file": "notes.md", "items": [i for x in md for i in _ids(x, "name", "phone")] +
                  _ids(cl[2], "first")})

    h = cl[7]
    (out_dir / "page.html").write_text(
        "<!doctype html><html><body><h1>Mijoz kartasi</h1><table>"
        f"<tr><td>F.I.Sh.</td><td>{h['name']}</td></tr>"
        f"<tr><td>Telefon</td><td><a href=\"tel:{h['phone']}\">{h['phone']}</a></td></tr>"
        f"<tr><td>Hisob</td><td><span data-account=\"{h['account']}\" title=\"{h['name']}\">{h['account']}</span>"
        "</td></tr></table></body></html>\n", encoding="utf-8")
    seeds.append({"file": "page.html", "items": _ids(h, "name", "phone", "account")})

    (out_dir / "seeds.json").write_text(json.dumps(seeds, ensure_ascii=False, indent=1), encoding="utf-8")
    return seeds


if __name__ == "__main__":
    build(Path(sys.argv[1] if len(sys.argv) > 1 else "eval/out/formats/files"))
